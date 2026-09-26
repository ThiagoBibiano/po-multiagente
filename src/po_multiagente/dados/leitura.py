"""Leitura de fontes CSV e XLSX, sem alterar os arquivos nem os valores.

A leitura interpreta a notação do texto (separador de campos, separador
decimal, datas), mas não transforma o dado: não agrega, não converte unidade,
não junta fontes nem preenche faltantes. Um valor que só seria número depois
de uma transformação, como ``"R$ 10,00"`` ou ``"12%"``, fica como texto e
leva a uma solicitação de tratamento.
"""

import csv
import io
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from po_multiagente.dados.erros import ErroDados
from po_multiagente.dominio import MotivoTratamento, TipoColuna

Celula = str | int | float | bool | date | None
"""Valor de uma célula depois da leitura; ``datetime`` é subtipo de ``date``."""

SEPARADORES_CAMPO = (";", ",", "\t", "|")
"""Candidatos a separador de campos, na ordem de desempate."""

_LINHAS_AMOSTRA_SEPARADOR = 50

_VERDADEIROS = frozenset({"true", "verdadeiro", "sim"})
_FALSOS = frozenset({"false", "falso", "não", "nao"})

_DATA_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATA_BR = re.compile(r"^\d{2}/\d{2}/\d{4}$")


@dataclass(frozen=True)
class _NotacaoNumerica:
    """Separadores decimal e de milhar de um arquivo CSV."""

    decimal: str
    milhar: str

    def inteiro(self, texto: str) -> int | None:
        """Lê um inteiro, com ou sem separador de milhar em grupos de três."""
        milhar = re.escape(self.milhar)
        if re.fullmatch(rf"[+-]?(\d{{1,3}}({milhar}\d{{3}})+|\d+)", texto) is None:
            return None
        return int(texto.replace(self.milhar, ""))

    def numero(self, texto: str) -> float | None:
        """Lê um número, com parte decimal opcional."""
        milhar, decimal = re.escape(self.milhar), re.escape(self.decimal)
        padrao = rf"[+-]?(\d{{1,3}}({milhar}\d{{3}})+|\d+)({decimal}\d+)?"
        if re.fullmatch(padrao, texto) is None:
            return None
        return float(texto.replace(self.milhar, "").replace(self.decimal, "."))


_NOTACAO_BR = _NotacaoNumerica(decimal=",", milhar=".")
_NOTACAO_US = _NotacaoNumerica(decimal=".", milhar=",")


@dataclass(frozen=True)
class Tabela:
    """Conteúdo de uma fonte já lida: colunas, tipos inferidos e linhas.

    Attributes:
        nome: Nome da fonte; o nome do arquivo ou, numa planilha com várias
            abas, ``arquivo.xlsx[aba]``.
        colunas: Nomes das colunas, como estão no cabeçalho.
        tipos: Tipo inferido de cada coluna, na mesma ordem.
        linhas: Valores por linha; linhas totalmente vazias são omitidas.
    """

    nome: str
    colunas: tuple[str, ...]
    tipos: tuple[TipoColuna, ...]
    linhas: tuple[tuple[Celula, ...], ...]

    def posicao(self, coluna: str) -> int:
        """Posição da coluna no cabeçalho.

        Raises:
            ErroDados: Se a coluna não existir na fonte.
        """
        try:
            return self.colunas.index(coluna)
        except ValueError:
            raise ErroDados(
                f"A fonte {self.nome!r} não tem a coluna {coluna!r}",
                arquivo=self.nome,
                coluna=coluna,
                motivo=MotivoTratamento.FALTANTE,
            ) from None

    def tipo(self, coluna: str) -> TipoColuna:
        """Tipo inferido da coluna."""
        return self.tipos[self.posicao(coluna)]

    def valores(self, coluna: str) -> tuple[Celula, ...]:
        """Valores da coluna, na ordem das linhas."""
        posicao = self.posicao(coluna)
        return tuple(linha[posicao] for linha in self.linhas)


def ler_csv(caminho: Path, *, nome: str | None = None) -> Tabela:
    """Lê um arquivo CSV.

    O separador de campos é o candidato que produz o mesmo número de colunas
    (maior que um) em todas as linhas da amostra. Com separador ``;``, a
    notação numérica é a brasileira (``1.234,56``); com ``,``, a
    norte-americana (``1,234.56``); com os demais, a que os próprios valores
    indicarem.

    Args:
        caminho: Arquivo a ler.
        nome: Nome da fonte; o padrão é o nome do arquivo.

    Raises:
        ErroDados: Se o arquivo estiver vazio, sem cabeçalho válido ou com
            linhas de tamanho diferente do cabeçalho.
    """
    nome = nome or caminho.name
    texto = _decodificar(caminho.read_bytes())
    separador = _detectar_separador(texto)
    linhas = [
        linha
        for linha in csv.reader(io.StringIO(texto, newline=""), delimiter=separador)
        if any(celula.strip() for celula in linha)
    ]
    if not linhas:
        raise ErroDados(f"A fonte {nome!r} está vazia", arquivo=nome)
    cabecalho, corpo = _validar_cabecalho(linhas[0], nome), linhas[1:]
    for numero, linha in enumerate(corpo, start=2):
        if len(linha) != len(cabecalho):
            raise ErroDados(
                f"A linha {numero} de {nome!r} tem {len(linha)} campos; o cabeçalho tem "
                f"{len(cabecalho)}",
                arquivo=nome,
            )
    notacao = _notacao_para(separador, corpo)
    colunas = [
        _inferir_coluna_texto([linha[i].strip() for linha in corpo], notacao)
        for i in range(len(cabecalho))
    ]
    return _montar_tabela(nome, cabecalho, colunas)


def ler_xlsx(caminho: Path) -> tuple[Tabela, ...]:
    """Lê as abas não vazias de uma planilha XLSX.

    Usa os valores gravados das células; fórmulas não são recalculadas. O
    tipo de cada coluna vem do tipo das células; um número digitado como
    texto continua texto.

    Returns:
        Uma tabela por aba não vazia. Com uma só aba, a tabela leva o nome do
        arquivo; com várias, ``arquivo.xlsx[aba]``.

    Raises:
        ErroDados: Se nenhuma aba tiver dados ou se um cabeçalho for inválido.
    """
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        abas = [
            (aba.title, [linha for linha in aba.iter_rows(values_only=True) if _tem_valor(linha)])
            for aba in livro.worksheets
        ]
    finally:
        livro.close()
    abas = [(titulo, linhas) for titulo, linhas in abas if linhas]
    if not abas:
        raise ErroDados(f"A fonte {caminho.name!r} está vazia", arquivo=caminho.name)
    tabelas = []
    for titulo, linhas in abas:
        nome = caminho.name if len(abas) == 1 else f"{caminho.name}[{titulo}]"
        cabecalho = _validar_cabecalho(["" if c is None else str(c) for c in linhas[0]], nome)
        corpo = [tuple(linha) + (None,) * (len(cabecalho) - len(linha)) for linha in linhas[1:]]
        excedentes = [
            linha for linha in corpo if any(c is not None for c in linha[len(cabecalho) :])
        ]
        if excedentes:
            raise ErroDados(f"A fonte {nome!r} tem valores fora das colunas nomeadas", arquivo=nome)
        colunas = [
            _inferir_coluna_nativa([_normalizar_celula(linha[i]) for linha in corpo])
            for i in range(len(cabecalho))
        ]
        tabelas.append(_montar_tabela(nome, cabecalho, colunas))
    return tuple(tabelas)


def _decodificar(conteudo: bytes) -> str:
    """Decodifica em UTF-8 (com ou sem BOM) ou, na falha, em Windows-1252."""
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        return conteudo.decode("cp1252")


def _detectar_separador(texto: str) -> str:
    amostra = texto.splitlines()[:_LINHAS_AMOSTRA_SEPARADOR]
    melhor, colunas_melhor = SEPARADORES_CAMPO[0], 1
    for candidato in SEPARADORES_CAMPO:
        tamanhos = {
            len(linha)
            for linha in csv.reader(amostra, delimiter=candidato)
            if any(c.strip() for c in linha)
        }
        if len(tamanhos) == 1 and (colunas := tamanhos.pop()) > colunas_melhor:
            melhor, colunas_melhor = candidato, colunas
    return melhor


def _notacao_para(separador: str, corpo: Sequence[Sequence[str]]) -> _NotacaoNumerica:
    if separador == ";":
        return _NOTACAO_BR
    if separador == ",":
        return _NOTACAO_US
    celulas = [c.strip() for linha in corpo for c in linha]
    virgula = any(re.fullmatch(r"[+-]?\d+,\d+", c) for c in celulas)
    ponto = any(re.fullmatch(r"[+-]?\d+\.\d+", c) for c in celulas)
    return _NOTACAO_BR if virgula and not ponto else _NOTACAO_US


def _validar_cabecalho(cabecalho: Sequence[str], nome: str) -> tuple[str, ...]:
    colunas = tuple(c.strip() for c in cabecalho)
    while colunas and not colunas[-1]:
        colunas = colunas[:-1]
    if not colunas or any(not c for c in colunas):
        raise ErroDados(f"A fonte {nome!r} tem coluna sem nome no cabeçalho", arquivo=nome)
    repetidas = sorted({c for c in colunas if colunas.count(c) > 1})
    if repetidas:
        raise ErroDados(
            f"A fonte {nome!r} tem colunas com nome repetido: {', '.join(repetidas)}",
            arquivo=nome,
        )
    return colunas


def _inferir_coluna_texto(
    textos: list[str], notacao: _NotacaoNumerica
) -> tuple[TipoColuna, list[Celula]]:
    """Infere o tipo de uma coluna de CSV e converte os valores para ele.

    O tipo é o primeiro, na ordem booleano, inteiro, decimal e data, que lê
    todos os valores preenchidos; se nenhum lê, a coluna é texto.
    """
    preenchidos = [t for t in textos if t]
    if not preenchidos:
        return TipoColuna.TEXTO, [None] * len(textos)
    leitores: list[tuple[TipoColuna, Callable[[str], Celula]]] = [
        (TipoColuna.BOOLEANO, _ler_booleano),
        (TipoColuna.DATA, _ler_data),
    ]
    # Zeros à esquerda indicam código, e não quantidade: "007" continua texto.
    if not any(re.fullmatch(r"[+-]?0\d+", t) for t in preenchidos):
        leitores[1:1] = [
            (TipoColuna.INTEIRO, notacao.inteiro),
            (TipoColuna.DECIMAL, notacao.numero),
        ]
    for tipo, leitor in leitores:
        convertidos = [leitor(t) if t else None for t in textos]
        if all(c is not None for c, t in zip(convertidos, textos, strict=True) if t):
            return tipo, convertidos
    return TipoColuna.TEXTO, [t or None for t in textos]


def _ler_booleano(texto: str) -> bool | None:
    minusculo = texto.lower()
    if minusculo in _VERDADEIROS:
        return True
    if minusculo in _FALSOS:
        return False
    return None


def _ler_data(texto: str) -> date | None:
    try:
        if _DATA_ISO.fullmatch(texto):
            return date.fromisoformat(texto)
        if _DATA_BR.fullmatch(texto):
            return datetime.strptime(texto, "%d/%m/%Y").date()
    except ValueError:
        return None
    return None


def _normalizar_celula(valor: object) -> Celula:
    """Converte o valor nativo do openpyxl num dos tipos de ``Celula``."""
    if valor is None or isinstance(valor, bool | int | float | date):
        return valor
    texto = str(valor).strip()
    return texto or None


def _inferir_coluna_nativa(valores: list[Celula]) -> tuple[TipoColuna, list[Celula]]:
    """Infere o tipo de uma coluna de planilha pelos tipos das células."""
    preenchidos = [v for v in valores if v is not None]
    if not preenchidos:
        return TipoColuna.TEXTO, valores
    if all(isinstance(v, bool) for v in preenchidos):
        return TipoColuna.BOOLEANO, valores
    if all(isinstance(v, int) and not isinstance(v, bool) for v in preenchidos):
        return TipoColuna.INTEIRO, valores
    if all(isinstance(v, int | float) and not isinstance(v, bool) for v in preenchidos):
        return TipoColuna.DECIMAL, valores
    if all(isinstance(v, date) for v in preenchidos):
        return TipoColuna.DATA, valores
    return TipoColuna.TEXTO, [None if v is None else texto_celula(v) for v in valores]


def texto_celula(valor: Celula) -> str:
    """Representação textual canônica de uma célula, usada como chave.

    Um inteiro gravado como ``1.0`` numa planilha vira ``"1"``, para que a
    mesma chave em duas fontes seja reconhecida como igual.
    """
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime) and valor.time() == datetime.min.time():
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor)


def _tem_valor(linha: Sequence[object]) -> bool:
    return any(c is not None and str(c).strip() for c in linha)


def _montar_tabela(
    nome: str,
    cabecalho: tuple[str, ...],
    colunas: list[tuple[TipoColuna, list[Celula]]],
) -> Tabela:
    return Tabela(
        nome=nome,
        colunas=cabecalho,
        tipos=tuple(tipo for tipo, _ in colunas),
        linhas=tuple(zip(*(valores for _, valores in colunas), strict=True)),
    )
